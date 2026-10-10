"""Bootstrap AHP from individual expert pairwise-comparison matrices."""
import re

import numpy as np
import pandas as pd
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt


RI = {1: 0.0, 2: 0.0, 3: 0.58, 4: 0.90, 5: 1.12, 6: 1.24,
      7: 1.32, 8: 1.41, 9: 1.45, 10: 1.49, 11: 1.51, 12: 1.48,
      13: 1.56, 14: 1.57, 15: 1.59}


def read_ahp_sheets(path, header_row=1, start_row=2, start_col=2):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = [s for s in wb.sheetnames if re.fullmatch(r"Ans\s*\d+", s, re.I)]
        sheets.sort(key=lambda s: int(re.search(r"\d+", s).group()))
        if len(sheets) < 2:
            raise ValueError("Bootstrap AHP requires at least two expert worksheets named Ans 1, Ans 2, ...")
        criteria = None
        matrices = []
        for sn in sheets:
            rows = list(wb[sn].iter_rows(values_only=True))
            if len(rows) < start_row:
                raise ValueError(f"Worksheet '{sn}' does not have enough rows")
            header = rows[header_row - 1]
            labels = []
            j = start_col - 1
            while j < len(header) and header[j] is not None and str(header[j]).strip():
                labels.append(str(header[j]).strip()); j += 1
            n = len(labels)
            if n < 2 or len(set(labels)) != n:
                raise ValueError(f"Header in '{sn}' must contain at least two unique criteria")
            if criteria is None: criteria = labels
            elif labels != criteria: raise ValueError(f"Criterion order in '{sn}' does not match the first worksheet")
            if len(rows) < start_row + n - 1:
                raise ValueError(f"Worksheet '{sn}' is missing rows from the AHP matrix")
            matrix = np.empty((n, n), dtype=float)
            for i in range(n):
                row = rows[start_row - 1 + i]
                row_label = row[start_col - 2] if start_col > 1 and len(row) >= start_col - 1 else None
                if row_label is not None and str(row_label).strip() and str(row_label).strip() != labels[i]:
                    raise ValueError(f"Row label '{row_label}' does not match '{labels[i]}' in '{sn}'")
                for j in range(n):
                    cell = row[start_col - 1 + j] if start_col - 1 + j < len(row) else None
                    try: matrix[i, j] = float(cell)
                    except (TypeError, ValueError): raise ValueError(f"AHP cell must be numeric in worksheet '{sn}', row {start_row+i}, column {start_col+j}")
            if np.any(~np.isfinite(matrix)) or np.any(matrix <= 0):
                raise ValueError(f"AHP matrix in '{sn}' must contain positive finite numbers")
            matrices.append(matrix)
        return criteria, matrices
    finally:
        wb.close()


def validate_reciprocal(matrix, tol=1e-3):
    m = np.asarray(matrix, dtype=float)
    if not np.allclose(np.diag(m), 1, atol=tol):
        raise ValueError("The AHP matrix diagonal must equal 1")
    if not np.allclose(m * m.T, 1, atol=tol, rtol=tol):
        raise ValueError("AHP matrix is not reciprocal: aᵢⱼ × aⱼᵢ must be approximately 1")


def ahp_weights(matrix):
    matrix = np.asarray(matrix, dtype=float)
    validate_reciprocal(matrix)
    vals, vecs = np.linalg.eig(matrix)
    idx = int(np.argmax(vals.real))
    lam = float(vals[idx].real)
    vector = np.abs(vecs[:, idx].real)
    weights = vector / vector.sum()
    n = len(weights)
    ci = max(0.0, (lam - n) / (n - 1)) if n > 1 else 0.0
    cr = ci / RI[n] if n in RI and RI[n] > 0 else 0.0
    return weights, lam, ci, cr


def aggregate_geometric(matrices):
    stack = np.asarray(matrices, dtype=float)
    return np.exp(np.mean(np.log(stack), axis=0))


def run_bootstrap_ahp(criteria, expert_matrices, B=2000, alpha=.05, seed=42):
    if len(expert_matrices) < 2:
        raise ValueError("Bootstrap AHP requires at least two experts")
    if B < 100 or not 0 < alpha < 1:
        raise ValueError("Use at least 100 bootstrap samples and set alpha between 0 and 1")
    n = len(criteria)
    if n > 15:
        raise ValueError("The random index table supports up to 15 criteria")
    if any(np.asarray(m).shape != (n, n) for m in expert_matrices):
        raise ValueError("AHP matrix dimensions do not match the number of criteria")
    for m in expert_matrices: validate_reciprocal(m)
    group = aggregate_geometric(expert_matrices)
    point, lam, ci, cr = ahp_weights(group)
    individual_cr = [ahp_weights(m)[3] for m in expert_matrices]
    rng = np.random.default_rng(seed)
    boot_w, boot_cr = np.empty((B, n)), np.empty(B)
    count = len(expert_matrices)
    for b in range(B):
        ids = rng.integers(0, count, count)
        boot_w[b], _, _, boot_cr[b] = ahp_weights(aggregate_geometric([expert_matrices[i] for i in ids]))
    lo, hi = alpha/2, 1-alpha/2
    rank_boot = np.argsort(np.argsort(-boot_w, axis=1), axis=1) + 1
    df = pd.DataFrame({
        "Criterion": criteria,
        "Group weight": point,
        "Weight mean (BS)": boot_w.mean(axis=0),
        "Weight SD (BS)": boot_w.std(axis=0, ddof=1),
        f"CI lower {int((1-alpha)*100)}%": np.quantile(boot_w, lo, axis=0),
        f"CI upper {int((1-alpha)*100)}%": np.quantile(boot_w, hi, axis=0),
        "Rank": pd.Series(point).rank(ascending=False, method="min").astype(int),
        "Mean rank (BS)": rank_boot.mean(axis=0),
        "Top rank frequency": (rank_boot == 1).mean(axis=0),
    }).sort_values("Rank").reset_index(drop=True)
    meta = {"group_lambda_max": lam, "group_ci": ci, "group_cr": cr,
            "individual_cr": individual_cr, "bootstrap_cr_mean": float(boot_cr.mean()),
            "bootstrap_cr_ci": [float(np.quantile(boot_cr, lo)), float(np.quantile(boot_cr, hi))]}
    return df, boot_w, meta


def export_ahp(df, boot_w, criteria, meta, output_xlsx, output_png):
    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Bootstrap AHP")
        pd.DataFrame({"Expert": [f"Ans {i+1}" for i in range(len(meta["individual_cr"]))],
                      "Individual CR": meta["individual_cr"]}).to_excel(writer, index=False, sheet_name="Expert consistency")
        pd.DataFrame(boot_w, columns=criteria).to_excel(writer, index=False, sheet_name="Bootstrap weights")
        pd.DataFrame([{"Group lambda max": meta["group_lambda_max"], "Group CI": meta["group_ci"],
                       "Group CR": meta["group_cr"], "Bootstrap CR mean": meta["bootstrap_cr_mean"],
                       "Bootstrap CR CI lower": meta["bootstrap_cr_ci"][0],
                       "Bootstrap CR CI upper": meta["bootstrap_cr_ci"][1]}]).to_excel(writer, index=False, sheet_name="Summary")
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                ws.column_dimensions[col[0].column_letter].width = min(34, max(12, max(len(str(c.value or "")) for c in col)+2))
    mean = df["Weight mean (BS)"].to_numpy()
    lo = next(c for c in df.columns if c.startswith("CI lower")); hi = next(c for c in df.columns if c.startswith("CI upper"))
    yerr = np.maximum(0, np.vstack([mean-df[lo].to_numpy(), df[hi].to_numpy()-mean]))
    fig, ax = plt.subplots(figsize=(max(8, len(df)*.8), 5.2))
    ax.bar(df["Criterion"], mean, color="#2E77D0", alpha=.85)
    ax.errorbar(df["Criterion"], mean, yerr=yerr, fmt="none", ecolor="#26364A", capsize=4)
    ax.set_ylabel("Priority weight"); ax.set_title("Bootstrap AHP: priority weights and percentile CI")
    ax.tick_params(axis="x", rotation=30); ax.grid(axis="y", alpha=.25)
    fig.tight_layout(); fig.savefig(output_png, dpi=180, bbox_inches="tight"); plt.close(fig)
