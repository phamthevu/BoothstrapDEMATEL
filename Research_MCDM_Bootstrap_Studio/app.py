"""Research MCDM Bootstrap Studio — native Tkinter desktop application."""
from __future__ import annotations

import os
import queue
import threading
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter

from bootstrap_dematel import (load_fuzzy_scale, build_defuzzify_func,
                               run_pipeline, validate_formula)
from bootstrap_topsis import (read_topsis_sheets, run_bootstrap_topsis,
                              export_topsis)
from bootstrap_ahp import read_ahp_sheets, run_bootstrap_ahp, export_ahp


BG = "#F3F6FA"
INK = "#172B4D"
BLUE = "#2457A7"
MUTED = "#607089"


class BootstrapApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Research MCDM Bootstrap Studio")
        self.geometry("1240x820")
        self.minsize(1080, 720)
        self.configure(bg=BG)
        self.input_path = tk.StringVar()
        self.fuzzy_path = tk.StringVar()
        self.output_dir = tk.StringVar(value=str(Path.cwd() / "outputs"))
        self.method = tk.StringVar(value="DEMATEL")
        self.formula = tk.StringVar(value="(l + m + u) / 3")
        self.B = tk.IntVar(value=2000)
        self.alpha = tk.DoubleVar(value=.05)
        self.seed = tk.IntVar(value=80)
        self.start_row = tk.IntVar(value=2)
        self.start_col = tk.IntVar(value=2)
        self.header_row = tk.IntVar(value=1)
        self.n_rows = tk.StringVar()
        self.n_cols = tk.StringVar()
        self.topsis_weights = tk.StringVar()
        self.topsis_types = tk.StringVar()
        self.preview_source = tk.StringVar(value="Expert workbook")
        self.preview_sheet = tk.StringVar()
        self.events = queue.Queue()
        self.latest_image = None
        self.latest_xlsx = None
        self._style()
        self._layout()
        self.method.trace_add("write", lambda *_: self._method_changed())
        self.after(150, self._poll_events)

    def _style(self):
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("TFrame", background=BG)
        style.configure("Card.TFrame", background="white")
        style.configure("TLabel", background=BG, foreground=INK, font=("Segoe UI", 10))
        style.configure("Card.TLabel", background="white", foreground=INK, font=("Segoe UI", 10))
        style.configure("Muted.TLabel", background="white", foreground=MUTED, font=("Segoe UI", 9))
        style.configure("Title.TLabel", background=BG, foreground=INK, font=("Segoe UI", 22, "bold"))
        style.configure("Sub.TLabel", background=BG, foreground=MUTED, font=("Segoe UI", 10))
        style.configure("Section.TLabel", background="white", foreground=BLUE, font=("Segoe UI", 12, "bold"))
        style.configure("TButton", font=("Segoe UI", 10), padding=(10, 7))
        style.configure("Accent.TButton", background=BLUE, foreground="white", font=("Segoe UI", 10, "bold"))
        style.map("Accent.TButton", background=[("active", "#194787"), ("disabled", "#9DB2D1")])
        style.configure("Treeview", rowheight=27, font=("Segoe UI", 9), background="white", fieldbackground="white")
        style.configure("Treeview.Heading", font=("Segoe UI", 9, "bold"), background="#E8EEF7", foreground=INK)
        style.configure("TNotebook", background=BG, borderwidth=0)
        style.configure("TNotebook.Tab", padding=(15, 8), font=("Segoe UI", 10, "bold"))

    def _layout(self):
        header = ttk.Frame(self, padding=(24, 20, 24, 12))
        header.pack(fill="x")
        ttk.Label(header, text="Research MCDM Bootstrap Studio", style="Title.TLabel").pack(anchor="w")
        ttk.Label(header, text="Bootstrap analysis for expert-based MCDM data", style="Sub.TLabel").pack(anchor="w", pady=(4, 0))
        body = ttk.Frame(self, padding=(22, 6, 22, 18))
        body.pack(fill="both", expand=True)
        self.tabs = ttk.Notebook(body)
        self.tabs.pack(fill="both", expand=True)
        setup = ttk.Frame(self.tabs, padding=16)
        preview = ttk.Frame(self.tabs, padding=16)
        results = ttk.Frame(self.tabs, padding=16)
        self.tabs.add(setup, text="  Analysis setup  ")
        self.tabs.add(preview, text="  Data preview  ")
        self.tabs.add(results, text="  Results  ")
        self._setup_tab(setup)
        self._preview_tab(preview)
        self._results_tab(results)

    def _card(self, parent, title, row, column, columnspan=1):
        card = ttk.Frame(parent, style="Card.TFrame", padding=16)
        card.grid(row=row, column=column, columnspan=columnspan, sticky="nsew", padx=8, pady=8)
        ttk.Label(card, text=title, style="Section.TLabel").pack(anchor="w", pady=(0, 10))
        # Keep the packed title separate from the card contents. Callers can
        # freely use grid inside this body without mixing geometry managers.
        body = ttk.Frame(card, style="Card.TFrame")
        body.pack(fill="both", expand=True)
        return body

    def _setup_tab(self, parent):
        parent.columnconfigure(0, weight=1)
        parent.columnconfigure(1, weight=1)
        parent.rowconfigure(1, weight=1)
        files = self._card(parent, "1  Data and output", 0, 0, 2)
        files_body = ttk.Frame(files, style="Card.TFrame")
        files_body.pack(fill="x")
        files_body.columnconfigure(1, weight=1)
        ttk.Label(files_body, text="Expert workbook", style="Card.TLabel").grid(row=0, column=0, sticky="w", pady=4)
        ttk.Entry(files_body, textvariable=self.input_path).grid(row=0, column=1, sticky="ew", padx=8)
        ttk.Button(files_body, text="Browse…", command=self._choose_input).grid(row=0, column=2)
        ttk.Button(files_body, text="Preview", command=lambda: self._open_preview("Expert workbook")).grid(row=0, column=3, padx=(6, 0))
        ttk.Label(files_body, text="Fuzzy scale (DEMATEL / fuzzy input)", style="Card.TLabel").grid(row=1, column=0, sticky="w", pady=4)
        ttk.Entry(files_body, textvariable=self.fuzzy_path).grid(row=1, column=1, sticky="ew", padx=8)
        ttk.Button(files_body, text="Browse…", command=self._choose_fuzzy).grid(row=1, column=2)
        ttk.Button(files_body, text="Preview", command=lambda: self._open_preview("Fuzzy scale")).grid(row=1, column=3, padx=(6, 0))
        ttk.Label(files_body, text="Save results to", style="Card.TLabel").grid(row=2, column=0, sticky="w", pady=4)
        ttk.Entry(files_body, textvariable=self.output_dir).grid(row=2, column=1, sticky="ew", padx=8)
        ttk.Button(files_body, text="Folder…", command=self._choose_output).grid(row=2, column=2)

        method = self._card(parent, "2  Method and bootstrap settings", 1, 0)
        method.columnconfigure(1, weight=1)
        ttk.Label(method, text="Method", style="Card.TLabel").grid(row=0, column=0, sticky="w", pady=5)
        ttk.Combobox(method, textvariable=self.method, values=["DEMATEL", "AHP", "TOPSIS"], state="readonly", width=22).grid(row=0, column=1, sticky="w", padx=6)
        ttk.Label(method, text="Bootstrap samples", style="Card.TLabel").grid(row=1, column=0, sticky="w", pady=5)
        ttk.Spinbox(method, from_=100, to=100000, increment=100, textvariable=self.B, width=12).grid(row=1, column=1, sticky="w", padx=6)
        ttk.Label(method, text="Confidence level", style="Card.TLabel").grid(row=2, column=0, sticky="w", pady=5)
        self.confidence = tk.StringVar(value="95%")
        ttk.Combobox(method, textvariable=self.confidence, values=["90%", "95%", "99%"], state="readonly", width=10).grid(row=2, column=1, sticky="w", padx=6)
        ttk.Label(method, text="Random seed", style="Card.TLabel").grid(row=3, column=0, sticky="w", pady=5)
        ttk.Spinbox(method, from_=0, to=2147483647, textvariable=self.seed, width=12).grid(row=3, column=1, sticky="w", padx=6)

        self.method_card = self._card(parent, "3  Method options", 1, 1)
        self.method_card.columnconfigure(1, weight=1)
        self.dematel_options = ttk.Frame(self.method_card, style="Card.TFrame")
        self.dematel_options.pack(fill="x")
        self.formula_widgets = self._field(self.dematel_options, "Defuzzification formula", self.formula, 0)
        self.formula_test_button = ttk.Button(self.dematel_options, text="Test", command=self._test_formula)
        self.formula_test_button.grid(row=0, column=2, padx=(4, 0))
        self.formula_help_button = ttk.Button(self.dematel_options, text="Help", command=self._formula_help)
        self.formula_help_button.grid(row=0, column=3, padx=(4, 0))
        ttk.Label(self.dematel_options, text="Header row / data row / data col", style="Card.TLabel").grid(row=1, column=0, sticky="w", pady=5)
        row = ttk.Frame(self.dematel_options, style="Card.TFrame"); row.grid(row=1, column=1, sticky="w")
        for var, width in [(self.header_row, 5), (self.start_row, 5), (self.start_col, 5)]:
            ttk.Spinbox(row, from_=1, to=10000, textvariable=var, width=width).pack(side="left", padx=(0, 6))
        ttk.Label(self.dematel_options, text="Factors (rows / cols; blank = auto)", style="Card.TLabel").grid(row=2, column=0, sticky="w", pady=5)
        nrow = ttk.Frame(self.dematel_options, style="Card.TFrame"); nrow.grid(row=2, column=1, sticky="w")
        ttk.Entry(nrow, textvariable=self.n_rows, width=7).pack(side="left", padx=(0, 5))
        ttk.Entry(nrow, textvariable=self.n_cols, width=7).pack(side="left")
        self.topsis_options = ttk.Frame(self.method_card, style="Card.TFrame")
        self._field(self.topsis_options, "Weights (optional, comma-separated)", self.topsis_weights, 0)
        self._field(self.topsis_options, "Types, in criterion order", self.topsis_types, 1)

        bottom = ttk.Frame(parent, padding=(8, 10, 8, 0))
        bottom.grid(row=2, column=0, columnspan=2, sticky="ew")
        self.run_button = ttk.Button(bottom, text="Run bootstrap analysis", style="Accent.TButton", command=self._run)
        self.run_button.pack(side="left")
        self.progress = ttk.Progressbar(bottom, mode="indeterminate", length=240)
        self.progress.pack(side="left", padx=14)
        self.status = tk.StringVar(value="Ready")
        ttk.Label(bottom, textvariable=self.status).pack(side="left", padx=6)

    def _field(self, parent, label, var, r):
        label_widget = ttk.Label(parent, text=label, style="Card.TLabel")
        entry_widget = ttk.Entry(parent, textvariable=var, width=38)
        label_widget.grid(row=r, column=0, sticky="w", pady=5)
        entry_widget.grid(row=r, column=1, sticky="ew", padx=6)
        return label_widget, entry_widget

    def _results_tab(self, parent):
        bar = ttk.Frame(parent)
        bar.pack(fill="x", pady=(0, 10))
        ttk.Label(bar, text="Analysis output", style="Section.TLabel").pack(side="left")
        ttk.Button(bar, text="Open chart", command=self._open_chart).pack(side="right", padx=4)
        ttk.Button(bar, text="Open Excel", command=self._open_excel).pack(side="right", padx=4)
        self.summary = tk.StringVar(value="")
        ttk.Label(parent, textvariable=self.summary, wraplength=1100, style="Sub.TLabel").pack(anchor="w", pady=(0, 10))
        frame = ttk.Frame(parent)
        frame.pack(fill="both", expand=True)
        self.tree = ttk.Treeview(frame, show="headings")
        vs = ttk.Scrollbar(frame, orient="vertical", command=self.tree.yview)
        hs = ttk.Scrollbar(frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=vs.set, xscrollcommand=hs.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        vs.grid(row=0, column=1, sticky="ns")
        hs.grid(row=1, column=0, sticky="ew")
        frame.rowconfigure(0, weight=1); frame.columnconfigure(0, weight=1)

    def _preview_tab(self, parent):
        toolbar = ttk.Frame(parent)
        toolbar.pack(fill="x", pady=(0, 10))
        ttk.Label(toolbar, text="Preview file", style="Section.TLabel").pack(side="left", padx=(0, 8))
        source_combo = ttk.Combobox(toolbar, textvariable=self.preview_source,
                                    values=["Expert workbook", "Fuzzy scale"], state="readonly", width=20,
                                    postcommand=self._refresh_preview)
        source_combo.pack(side="left")
        source_combo.bind("<<ComboboxSelected>>", lambda _event: self._refresh_preview())
        ttk.Label(toolbar, text="Worksheet", padding=(14, 0, 6, 0)).pack(side="left")
        self.preview_sheet_combo = ttk.Combobox(toolbar, textvariable=self.preview_sheet,
                                                state="readonly", width=28)
        self.preview_sheet_combo.pack(side="left")
        self.preview_sheet_combo.bind("<<ComboboxSelected>>", lambda _event: self._render_preview())
        ttk.Button(toolbar, text="Refresh", command=self._refresh_preview).pack(side="left", padx=8)
        self.preview_status = tk.StringVar(value="Select a workbook or fuzzy scale, then preview its worksheets here.")
        ttk.Label(parent, textvariable=self.preview_status, style="Sub.TLabel").pack(anchor="w", pady=(0, 8))
        table_frame = ttk.Frame(parent)
        table_frame.pack(fill="both", expand=True)
        self.preview_tree = ttk.Treeview(table_frame, show="headings")
        yscroll = ttk.Scrollbar(table_frame, orient="vertical", command=self.preview_tree.yview)
        xscroll = ttk.Scrollbar(table_frame, orient="horizontal", command=self.preview_tree.xview)
        self.preview_tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        self.preview_tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        xscroll.grid(row=1, column=0, sticky="ew")
        table_frame.rowconfigure(0, weight=1)
        table_frame.columnconfigure(0, weight=1)

    def _method_changed(self):
        if not hasattr(self, "dematel_options"):
            return
        # Reflow the method panel without allowing the toplevel window to jump
        # to a different requested size when its controls change.
        current_geometry = self.geometry()
        current_state = self.state()
        self.dematel_options.pack_forget(); self.topsis_options.pack_forget()
        if self.method.get() in {"DEMATEL", "AHP"}:
            self.dematel_options.pack(fill="x")
        else:
            self.topsis_options.pack(fill="x")
        show_formula = self.method.get() == "DEMATEL"
        for widget in self.formula_widgets:
            if show_formula: widget.grid()
            else: widget.grid_remove()
        def restore_window(size=current_geometry, state=current_state):
            if state == "zoomed":
                self.state("zoomed")
            else:
                self.geometry(size)
        self.after_idle(restore_window)

    def _formula_help(self):
        messagebox.showinfo(
            "Defuzzification formula",
            "Use the variables l, m, and u with arithmetic operators +, -, *, /, **, and %.\n\n"
            "Examples:\n"
            "(l + m + u) / 3\n"
            "(l + 2*m + u) / 4\n"
            "m**0.5 for the square root of m\n\n"
            "Do not use ^ for powers; use **."
        )

    def _test_formula(self):
        valid, error = validate_formula(self.formula.get())
        if not valid:
            messagebox.showerror("Formula check", error)
            return
        value = build_defuzzify_func(self.formula.get())(1, 2, 3)
        messagebox.showinfo("Formula check", f"Formula is valid.\n\nExample result for l=1, m=2, u=3: {value:g}")

    def _open_preview(self, source):
        self.preview_source.set(source)
        self._refresh_preview()
        self.tabs.select(1)

    def _preview_path(self):
        return self.input_path.get() if self.preview_source.get() == "Expert workbook" else self.fuzzy_path.get()

    def _refresh_preview(self, *_args):
        path = self._preview_path()
        if not path or not Path(path).is_file():
            source = self.preview_source.get().lower()
            self.preview_sheet_combo.configure(values=[])
            self.preview_sheet.set("")
            self._clear_preview()
            self.preview_status.set(f"No {source} selected. Choose its Excel file on the Analysis setup tab.")
            return
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            sheet_names = wb.sheetnames
            wb.close()
            self.preview_sheet_combo.configure(values=sheet_names)
            if self.preview_sheet.get() not in sheet_names:
                self.preview_sheet.set(sheet_names[0])
            self._render_preview()
        except Exception as exc:
            self._clear_preview()
            self.preview_status.set(f"Could not preview this workbook: {exc}")

    def _clear_preview(self):
        self.preview_tree.delete(*self.preview_tree.get_children())
        self.preview_tree.configure(columns=())

    def _render_preview(self):
        path = self._preview_path()
        sheet_name = self.preview_sheet.get()
        if not path or not Path(path).is_file() or not sheet_name:
            return
        wb = None
        try:
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb[sheet_name]
            max_rows = ws.max_row or 0
            max_columns = ws.max_column or 0
            row_limit, column_limit = 100, 40
            is_fuzzy_scale = self.preview_source.get() == "Fuzzy scale"
            if is_fuzzy_scale:
                # Show the parsed scale used by the analysis, matching the old
                # app's fuzzy-label preview instead of only showing raw cells.
                scale = load_fuzzy_scale(path)
                headers = ["Label", "L", "M", "U"]
                rows = [[label, *values] for label, values in scale.items()]
                shown_rows = min(len(rows), row_limit)
                shown_columns = len(headers)
                rows = rows[:row_limit]
                columns = headers
            else:
                shown_rows, shown_columns = min(max_rows, row_limit), min(max_columns, column_limit)
                columns = [get_column_letter(i) for i in range(1, shown_columns + 1)]
                rows = [list(values) for values in ws.iter_rows(
                    min_row=1, max_row=shown_rows, max_col=shown_columns, values_only=True)]
            self.preview_tree.delete(*self.preview_tree.get_children())
            self.preview_tree.configure(columns=columns)
            for column in columns:
                self.preview_tree.heading(column, text=column)
                self.preview_tree.column(column, width=180 if column == "Label" else 125,
                                         minwidth=80, anchor="center", stretch=True)
            for values in rows:
                display = ["" if value is None else str(value) for value in values]
                self.preview_tree.insert("", "end", values=display)
            if is_fuzzy_scale:
                self.preview_status.set(
                    f"{Path(path).name}  ·  {len(scale):,} fuzzy labels loaded  ·  "
                    f"Showing {shown_rows:,} labels (L, M, U)."
                )
            else:
                self.preview_status.set(
                    f"{Path(path).name}  ·  {sheet_name}  ·  Showing {shown_rows:,} of {max_rows:,} rows and "
                    f"{shown_columns:,} of {max_columns:,} columns."
                )
        except Exception as exc:
            self._clear_preview()
            self.preview_status.set(f"Could not preview worksheet '{sheet_name}': {exc}")
        finally:
            if wb is not None:
                wb.close()

    def _choose_input(self):
        p = filedialog.askopenfilename(filetypes=[("Excel workbooks", "*.xlsx *.xlsm")])
        if p:
            self.input_path.set(p)
            self._open_preview("Expert workbook")
    def _choose_fuzzy(self):
        p = filedialog.askopenfilename(filetypes=[("Excel workbooks", "*.xlsx")])
        if p:
            self.fuzzy_path.set(p)
            self._open_preview("Fuzzy scale")
    def _choose_output(self):
        p = filedialog.askdirectory()
        if p: self.output_dir.set(p)

    def _run(self):
        if not self.input_path.get() or not Path(self.input_path.get()).is_file():
            messagebox.showerror("Workbook required", "Select a readable Excel workbook first."); return
        if self.method.get() == "DEMATEL" and (not self.fuzzy_path.get() or not Path(self.fuzzy_path.get()).is_file()):
            messagebox.showerror("Fuzzy scale required", "Select an Excel fuzzy scale for DEMATEL."); return
        if self.B.get() < 100:
            messagebox.showerror("Invalid bootstrap count", "Use at least 100 bootstrap samples."); return
        self.run_button.configure(state="disabled")
        self.progress.start(12)
        self.status.set("Running analysis…")
        self.tabs.select(0)
        threading.Thread(target=self._worker, daemon=True).start()

    def _worker(self):
        try:
            out_dir = Path(self.output_dir.get()).expanduser()
            out_dir.mkdir(parents=True, exist_ok=True)
            stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            prefix = f"{self.method.get().lower()}_bootstrap_{stamp}"
            xlsx, png = out_dir / f"{prefix}.xlsx", out_dir / f"{prefix}.png"
            alpha = 1 - int(self.confidence.get().strip("%")) / 100
            if self.method.get() == "DEMATEL":
                valid, error = validate_formula(self.formula.get())
                if not valid: raise ValueError(f"Defuzzification formula: {error}")
                from bootstrap_dematel import load_fuzzy_scale, read_expert_sheets, run_bootstrap, export_excel, plot_irm
                scale = load_fuzzy_scale(self.fuzzy_path.get())
                factors, experts = read_expert_sheets(self.input_path.get(), start_row=self.start_row.get(),
                    start_col=self.start_col.get(), n_rows=int(self.n_rows.get()) if self.n_rows.get().strip() else None,
                    n_cols=int(self.n_cols.get()) if self.n_cols.get().strip() else None,
                    header_row=self.header_row.get(), fuzzy_scale=scale,
                    defuzz_func=build_defuzzify_func(self.formula.get()))
                df, _, _, T, pb, rb = run_bootstrap(factors, experts, B=self.B.get(), seed=self.seed.get(), alpha=alpha)
                export_excel(df, factors, T, pb, rb, str(xlsx), self.B.get(), alpha)
                plot_irm(df, factors, pb, rb, str(png), alpha)
                summary = f"DEMATEL · {len(experts)} experts · {len(factors)} factors · {self.B.get():,} resamples · {self.confidence.get()} percentile CI. Direction stability is descriptive and is not a p-value."
            elif self.method.get() == "TOPSIS":
                scale = load_fuzzy_scale(self.fuzzy_path.get()) if self.fuzzy_path.get().strip() else None
                alternatives, criteria, matrices = read_topsis_sheets(
                    self.input_path.get(), fuzzy_scale=scale,
                    defuzz_func=build_defuzzify_func(self.formula.get()) if scale else None)
                weights = [float(x.strip()) for x in self.topsis_weights.get().split(",")] if self.topsis_weights.get().strip() else None
                types = [x.strip().lower() for x in self.topsis_types.get().split(",")] if self.topsis_types.get().strip() else ["benefit"] * len(criteria)
                if len(types) != len(criteria) or any(x not in {"benefit", "cost"} for x in types):
                    raise ValueError(f"Enter exactly {len(criteria)} criterion types using Benefit or Cost")
                df, boot = run_bootstrap_topsis(alternatives, criteria, matrices, B=self.B.get(), alpha=alpha,
                    seed=self.seed.get(), weights=weights, benefit=[x == "benefit" for x in types])
                weights_norm = (pd.Series(weights if weights is not None else [1]*len(criteria), dtype=float) /
                                sum(weights if weights is not None else [1]*len(criteria))).tolist()
                export_topsis(df, boot, criteria, weights_norm, alternatives, str(xlsx), str(png))
                summary = f"TOPSIS · {len(matrices)} experts · {len(alternatives)} alternatives · {len(criteria)} criteria · {self.B.get():,} resamples · {self.confidence.get()} percentile CI."
            else:
                criteria, matrices = read_ahp_sheets(self.input_path.get(), header_row=self.header_row.get(),
                    start_row=self.start_row.get(), start_col=self.start_col.get())
                df, boot, meta = run_bootstrap_ahp(criteria, matrices, B=self.B.get(), alpha=alpha, seed=self.seed.get())
                export_ahp(df, boot, criteria, meta, str(xlsx), str(png))
                cr_status = "within 0.10 guideline" if meta["group_cr"] <= .10 else "above 0.10 guideline"
                summary = f"AHP · {len(matrices)} experts · {len(criteria)} criteria · {self.B.get():,} resamples · group CR={meta['group_cr']:.4f} ({cr_status}). Individual CRs and bootstrap CR interval are in the workbook."
            self.events.put(("done", df, summary, str(xlsx), str(png)))
        except Exception as exc:
            self.events.put(("error", str(exc)))

    def _poll_events(self):
        try:
            while True:
                event = self.events.get_nowait()
                self.progress.stop(); self.run_button.configure(state="normal")
                if event[0] == "error":
                    self.status.set("Analysis stopped with an error.")
                    messagebox.showerror("Analysis failed", event[1])
                else:
                    _, df, summary, self.latest_xlsx, self.latest_image = event
                    self._show_results(df, summary)
                    self.status.set("Analysis complete.")
                    self.tabs.select(1)
        except queue.Empty:
            pass
        self.after(150, self._poll_events)

    def _show_results(self, df, summary):
        self.summary.set(summary)
        self.tree.delete(*self.tree.get_children())
        columns = list(df.columns)
        self.tree.configure(columns=columns)
        for col in columns:
            self.tree.heading(col, text=col)
            self.tree.column(col, width=min(210, max(100, len(col)*9)), anchor="center", stretch=True)
        for _, row in df.iterrows():
            vals = [f"{v:.4f}" if isinstance(v, float) else str(v) for v in row.tolist()]
            self.tree.insert("", "end", values=vals)

    def _open_chart(self):
        self._open(self.latest_image)
    def _open_excel(self):
        self._open(self.latest_xlsx)
    @staticmethod
    def _open(path):
        if not path or not Path(path).exists():
            messagebox.showinfo("No result yet", "Run an analysis first."); return
        try:
            if os.name == "nt": os.startfile(path)
            elif os.name == "posix":
                import subprocess
                subprocess.Popen(["open" if __import__("sys").platform == "darwin" else "xdg-open", path])
        except Exception as e:
            messagebox.showerror("Could not open file", str(e))


if __name__ == "__main__":
    BootstrapApp().mainloop()
