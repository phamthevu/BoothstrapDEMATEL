"""Bootstrap analysis for crisp or Z-fuzzy DEMATEL expert data."""

import numpy as np
import pandas as pd
import openpyxl
from openpyxl.styles import (PatternFill, Font, Alignment, Border, Side,
                              GradientFill)
from openpyxl.utils import get_column_letter
from openpyxl.chart import ScatterChart, Reference, Series
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyArrowPatch
import re
import sys
import ast
import warnings
warnings.filterwarnings('ignore')

# Z-fuzzy scale (L, M, U)
def load_fuzzy_scale(filepath):
    """Load fuzzy labels from Excel columns A:D: label, L, M, U."""
    wb = openpyxl.load_workbook(filepath, data_only=True)
    ws = wb.active

    scale = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] is None:
            continue

        key = str(row[0]).strip()
        try:
            l = float(row[1])
            m = float(row[2])
            u = float(row[3])
        except:
            continue

        scale[key] = (l, m, u)

    wb.close()

    if "0" not in scale:
        scale["0"] = (0.0, 0.0, 0.0)

    return scale

def parse_fuzzy_label(val, fuzzy_scale):
    if val is None or val == 0 or str(val).strip() == "0":
        return (0.0, 0.0, 0.0)

    key = str(val).strip()

    if key in fuzzy_scale:
        return fuzzy_scale[key]

    raise ValueError(f"Fuzzy label not found: {key}")

# Custom defuzzification formula
def build_defuzzify_func(formula_str):
    """Build a safe numeric function from an expression using l, m, and u."""
    if formula_str is None or formula_str.strip() == "":
        # Use the arithmetic mean by default.
        return lambda l, m, u: (l + m + u) / 3

    tree = ast.parse(formula_str, mode="eval")
    allowed = (ast.Expression, ast.BinOp, ast.UnaryOp, ast.Add, ast.Sub,
               ast.Mult, ast.Div, ast.Pow, ast.USub, ast.UAdd, ast.Name,
               ast.Constant, ast.Load, ast.Mod)
    for node in ast.walk(tree):
        if not isinstance(node, allowed):
            raise ValueError("Formula may only use l, m, u and the operators + - * / ** %")
        if isinstance(node, ast.Name) and node.id not in {"l", "m", "u"}:
            raise ValueError(f"Invalid variable: {node.id}")
        if isinstance(node, ast.Constant) and not isinstance(node.value, (int, float)):
            raise ValueError("Only numeric constants are allowed")
    code = compile(tree, "<defuzzify>", "eval")

    def func(l, m, u):
        result = float(eval(code, {"__builtins__": {}}, {"l": l, "m": m, "u": u}))
        if not np.isfinite(result):
            raise ValueError("The formula returned a non-finite value")
        return result

    return func

def validate_formula(formula_str):
    try:
        f = build_defuzzify_func(formula_str)
        test = f(1, 2, 3)
        return True, ""
    except Exception as e:
        return False, str(e)

# Read expert worksheets
def read_expert_sheets(
    filepath,
    start_row=2,
    start_col=2,
    n_rows=None,
    n_cols=None,
    header_row=1,
    fuzzy_scale=None,
    defuzz_func=None
):
    """
    Configurable reader:
    - start_row: first data row (1-based)
    - start_col: first data column (1-based)
    - n_rows: number of factor rows
    - n_cols: number of factor columns
    - header_row: row containing factor labels
    """

    wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
    sheet_names = wb.sheetnames

    ans_sheets = [s for s in sheet_names if re.match(r'^Ans\s*\d+$', s, re.IGNORECASE)]
    ans_sheets.sort(key=lambda x: int(re.search(r'\d+', x).group()))

    if not ans_sheets:
        wb.close()
        raise ValueError("No expert worksheets found. Name them Ans 1, Ans 2, ...")

    # Read factor headers.
    ws0 = wb[ans_sheets[0]]
    rows0 = list(ws0.iter_rows(values_only=True))

    header = rows0[header_row - 1]

    if n_cols is None:
        # Detect the matrix width from the first blank header cell.
        factors = []
        for c in header[start_col - 1:]:
            if c is None:
                break
            factors.append(str(c).strip())
        n = len(factors)
    else:
        factors = [str(header[start_col - 1 + i]).strip() for i in range(n_cols)]
        n = n_cols

    if len(set(factors)) != len(factors):
        wb.close()
        raise ValueError("Factor names in the header must be unique")

    if n_rows is not None and int(n_rows) != len(factors):
        wb.close()
        raise ValueError(f"Row count ({n_rows}) must match the number of factors in the header ({len(factors)})")
    if n_cols is not None and int(n_cols) != len(factors):
        wb.close()
        raise ValueError(f"Column count ({n_cols}) must match the number of factors in the header ({len(factors)})")
    if n < 2:
        wb.close()
        raise ValueError("DEMATEL requires at least two factors")

    # ===== Read expert matrices =====
    experts = []

    for sheet_name in ans_sheets:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))

        sheet_header = rows[header_row - 1]
        sheet_factors = [str(sheet_header[start_col - 1 + i]).strip() for i in range(n)]
        if sheet_factors != factors:
            wb.close()
            raise ValueError(f"Factor order in worksheet '{sheet_name}' does not match the first expert worksheet")

        matrix = np.zeros((n, n))

        if len(rows) < start_row + n - 1:
            wb.close()
            raise ValueError(f"Worksheet '{sheet_name}' does not contain enough rows for a {n}×{n} matrix")
        for i in range(n):
            if start_col > 1:
                row_label = rows[start_row - 1 + i][start_col - 2]
                if row_label is not None and str(row_label).strip() and str(row_label).strip() != factors[i]:
                    wb.close()
                    raise ValueError(f"Row label '{row_label}' in worksheet '{sheet_name}' does not match factor '{factors[i]}'")
            for j in range(n):
                try:
                    val = rows[start_row - 1 + i][start_col - 1 + j]
                    l, m, u = parse_fuzzy_label(val, fuzzy_scale)
                    matrix[i, j] = defuzz_func(l, m, u)
                except Exception as exc:
                    wb.close()
                    raise ValueError(f"Error in worksheet '{sheet_name}', row {start_row+i}, column {start_col+j}: {exc}") from exc
        np.fill_diagonal(matrix, 0.0)

        experts.append(matrix)

    wb.close()

    print(f"✅ Read {len(experts)} experts, {n} factors")
    return factors, experts

# Calculate DEMATEL from a set of expert matrices.
def compute_dematel(matrices):
    """
    matrices: list of n×n NumPy arrays
    Returns r+c (prominence), r-c (relation), r, c, and T (total influence matrix).
    """
    if not matrices:
        raise ValueError("No expert matrices were provided")
    n = matrices[0].shape[0]
    if any(np.asarray(m).shape != (n, n) for m in matrices):
        raise ValueError("All expert matrices must be square and have the same dimensions")
    # Average the matrices.
    A = np.mean(matrices, axis=0)

    # Normalize the direct-relation matrix.
    row_sums = A.sum(axis=1)
    col_sums = A.sum(axis=0)
    k = max(row_sums.max(), col_sums.max())
    if k == 0:
        return np.zeros(n), np.zeros(n), np.zeros(n), np.zeros(n), A
    D = A / k

    # Total influence matrix: T = D(I-D)^-1.
    I = np.eye(n)
    try:
        T = np.linalg.solve(I - D, D)
    except np.linalg.LinAlgError as exc:
        raise ValueError("Could not normalize the DEMATEL matrix: I-D is singular") from exc

    r = T.sum(axis=1)  # Row sums.
    c = T.sum(axis=0)  # Column sums.

    return r + c, r - c, r, c, T

# Bootstrap analysis.
def run_bootstrap(factors, experts, B=2000, seed=42, alpha=0.05):
    """
    Run B bootstrap resamples and calculate intervals for r+c and r-c.
    """
    if len(experts) < 2:
        raise ValueError("Bootstrap requires at least two experts to estimate uncertainty")
    if len(factors) < 2 or any(np.asarray(m).shape != (len(factors), len(factors)) for m in experts):
        raise ValueError("The number of factors does not match the expert matrix dimensions")
    if B < 100:
        raise ValueError("Use at least 100 bootstrap samples")
    if not 0 < alpha < 1:
        raise ValueError("Alpha must be between 0 and 1")
    rng = np.random.default_rng(seed)
    n_exp = len(experts)
    n_fac = len(factors)

    prominence_boot = np.zeros((B, n_fac))
    relation_boot   = np.zeros((B, n_fac))

    for b in range(B):
        idx = rng.integers(0, n_exp, size=n_exp)
        sample = [experts[i] for i in idx]
        rc_plus, rc_minus, _, _, _ = compute_dematel(sample)
        prominence_boot[b] = rc_plus
        relation_boot[b]   = rc_minus

    # Original estimates without resampling.
    rc_plus_orig, rc_minus_orig, r_orig, c_orig, T_orig = compute_dematel(experts)
    original_ranks = pd.Series(rc_plus_orig).rank(ascending=False, method="min").astype(int).to_numpy()

    lo = alpha / 2
    hi = 1 - alpha / 2

    results = []
    for i, fac in enumerate(factors):
        p_mean  = prominence_boot[:, i].mean()
        p_std   = prominence_boot[:, i].std()
        p_lo    = np.percentile(prominence_boot[:, i], lo*100)
        p_hi    = np.percentile(prominence_boot[:, i], hi*100)

        r_mean  = relation_boot[:, i].mean()
        r_std   = relation_boot[:, i].std()
        r_lo    = np.percentile(relation_boot[:, i], lo*100)
        r_hi    = np.percentile(relation_boot[:, i], hi*100)

        # Descriptive direction stability, not a hypothesis-test p-value.
        orig_sign = np.sign(rc_minus_orig[i])
        if orig_sign > 0:
            stability = (relation_boot[:, i] > 0).mean()
        elif orig_sign < 0:
            stability = (relation_boot[:, i] < 0).mean()
        else:
            stability = 0.0

        # A direction is stable when its interval excludes zero.
        sig = "Stable" if (r_lo > 0 or r_hi < 0) else "Uncertain"
        cause_effect = "Cause" if rc_minus_orig[i] > 0 else "Effect" if rc_minus_orig[i] < 0 else "Neutral"

        # Normalized prominence weight.
        total_prominence = rc_plus_orig.sum()
        w = rc_plus_orig[i] / total_prominence if total_prominence else 0.0

        results.append({
            "Factor": fac,
            "r (original)": round(r_orig[i], 4),
            "c (original)": round(c_orig[i], 4),
            "r+c (original)": round(rc_plus_orig[i], 4),
            "r-c (original)": round(rc_minus_orig[i], 4),
            "Cause/Effect": cause_effect,
            "Weight": round(w, 4),
            "Rank": int(original_ranks[i]),
            # Bootstrap r+c.
            "Bootstrap mean (r+c)": round(p_mean, 4),
            "Bootstrap SD (r+c)":   round(p_std, 4),
            f"Bootstrap CI lower (r+c) {int((1-alpha)*100)}%": round(p_lo, 4),
            f"Bootstrap CI upper (r+c) {int((1-alpha)*100)}%": round(p_hi, 4),
            # Bootstrap r-c.
            "Bootstrap mean (r-c)": round(r_mean, 4),
            "Bootstrap SD (r-c)":   round(r_std, 4),
            f"Bootstrap CI lower (r-c) {int((1-alpha)*100)}%": round(r_lo, 4),
            f"Bootstrap CI upper (r-c) {int((1-alpha)*100)}%": round(r_hi, 4),
            # Direction stability summary.
            "Direction stability": round(stability, 4),
            "CI excludes zero?": sig,
        })

    # Sort by descending original prominence rank.
    df = pd.DataFrame(results)
    df = df.sort_values("Rank")

    return df, rc_plus_orig, rc_minus_orig, T_orig, prominence_boot, relation_boot

# Export formatted Excel results.
def style_header(ws, row, cols, fill_hex, font_color="FFFFFF", font_size=11):
    fill = PatternFill("solid", fgColor=fill_hex)
    font = Font(bold=True, color=font_color, size=font_size, name="Arial")
    for col in range(1, cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def style_data_row(ws, row, n_cols, is_alt):
    fill = PatternFill("solid", fgColor="F0F4FF" if is_alt else "FFFFFF")
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.font = Font(name="Arial", size=10)

def add_thin_border(ws, min_row, max_row, min_col, max_col):
    thin = Side(style='thin', color="CCCCCC")
    for row in ws.iter_rows(min_row=min_row, max_row=max_row,
                             min_col=min_col, max_col=max_col):
        for cell in row:
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)

def export_excel(df, factors, T_orig, prominence_boot, relation_boot,
                 output_path, B, alpha):
    wb = openpyxl.Workbook()

    # Sheet 1: bootstrap results.
    ws1 = wb.active
    ws1.title = "DEMATEL Bootstrap"
    ws1.sheet_view.showGridLines = False

    # Main title.
    ws1.merge_cells("A1:V1")
    title_cell = ws1["A1"]
    title_cell.value = f"Z-Fuzzy DEMATEL Bootstrap Results (B={B:,}, CI={(1-alpha)*100:.0f}%)"
    title_cell.font = Font(bold=True, size=14, color="FFFFFF", name="Arial")
    title_cell.fill = PatternFill("solid", fgColor="1A3A6B")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws1.row_dimensions[1].height = 32

    # Header
    headers = list(df.columns)
    n_cols = len(headers)
    for col_idx, h in enumerate(headers, 1):
        ws1.cell(row=2, column=col_idx, value=h)
    style_header(ws1, 2, n_cols, "2E5FA3")
    ws1.row_dimensions[2].height = 28

    # Color cause/effect labels.
    cause_fill = PatternFill("solid", fgColor="DFF0D8")
    effect_fill = PatternFill("solid", fgColor="F2DEDE")
    sig_fill = PatternFill("solid", fgColor="FFF3CD")

    for row_idx, (_, row_data) in enumerate(df.iterrows(), 3):
        is_alt = (row_idx % 2 == 0)
        style_data_row(ws1, row_idx, n_cols, is_alt)
        for col_idx, val in enumerate(row_data, 1):
            cell = ws1.cell(row=row_idx, column=col_idx, value=val)
            # Highlight cause/effect labels.
            if headers[col_idx - 1] == "Cause/Effect":
                cell.fill = cause_fill if val == "Cause" else effect_fill
                cell.font = Font(bold=True, name="Arial", size=10,
                                 color="2D6A2D" if val == "Cause" else "A94442")
            if headers[col_idx - 1] == "CI excludes zero?":
                if val == "Stable":
                    cell.fill = sig_fill
                    cell.font = Font(bold=True, name="Arial", size=10, color="8A6D3B")

    add_thin_border(ws1, 2, 2 + len(df), 1, n_cols)

    # Column widths.
    col_widths = [12, 10, 10, 12, 12, 13, 12, 10,
                  14, 12, 18, 18, 14, 12, 18, 18,
                  10, 18]
    for i, w in enumerate(col_widths[:n_cols], 1):
        ws1.column_dimensions[get_column_letter(i)].width = w

    # Statistical note.
    note_row = 3 + len(df) + 1
    ws1.cell(row=note_row, column=1,
             value=f"Note: B={B:,} bootstrap samples | CI={int((1-alpha)*100)}% percentile bootstrap | "
                   f"Direction stability is the share of resamples matching the original sign; it is not a p-value.").font = Font(
        italic=True, size=9, color="555555", name="Arial")

    # Sheet 2: total influence matrix.
    ws2 = wb.create_sheet("Total Influence Matrix")
    ws2.sheet_view.showGridLines = False

    ws2.merge_cells(f"A1:{get_column_letter(len(factors)+1)}1")
    t2 = ws2["A1"]
    t2.value = "TOTAL INFLUENCE MATRIX (T)"
    t2.font = Font(bold=True, size=13, color="FFFFFF", name="Arial")
    t2.fill = PatternFill("solid", fgColor="1A3A6B")
    t2.alignment = Alignment(horizontal="center", vertical="center")

    ws2.cell(row=2, column=1, value="")
    for j, fac in enumerate(factors, 2):
        ws2.cell(row=2, column=j, value=fac)
    style_header(ws2, 2, len(factors)+1, "2E5FA3")

    max_t = T_orig.max()
    for i, fac in enumerate(factors):
        ws2.cell(row=i+3, column=1, value=fac).font = Font(bold=True, name="Arial", size=10)
        ws2.cell(row=i+3, column=1).fill = PatternFill("solid", fgColor="2E5FA3")
        ws2.cell(row=i+3, column=1).font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
        for j in range(len(factors)):
            val = T_orig[i, j]
            cell = ws2.cell(row=i+3, column=j+2, value=round(val, 4))
            # Blue heatmap.
            intensity = int(255 - (val / max_t) * 180) if max_t > 0 else 255
            hex_col = f"{'%02X' % intensity}{'%02X' % intensity}FF"
            cell.fill = PatternFill("solid", fgColor=hex_col)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(name="Arial", size=9)

    add_thin_border(ws2, 2, 2+len(factors), 1, len(factors)+1)
    for col in range(1, len(factors)+2):
        ws2.column_dimensions[get_column_letter(col)].width = 10

    # Sheet 3: bootstrap distribution summary.
    ws3 = wb.create_sheet("Bootstrap Distribution")
    ws3.sheet_view.showGridLines = False

    ws3.merge_cells("A1:F1")
    d3 = ws3["A1"]
    d3.value = "BOOTSTRAP DISTRIBUTION SUMMARY (r-c)"
    d3.font = Font(bold=True, size=13, color="FFFFFF", name="Arial")
    d3.fill = PatternFill("solid", fgColor="1A3A6B")
    d3.alignment = Alignment(horizontal="center", vertical="center")

    bs_headers = ["Factor", "Min", "P10", "P25", "Median", "P75", "P90", "Max", "Mean", "SD"]
    for ci, h in enumerate(bs_headers, 1):
        ws3.cell(row=2, column=ci, value=h)
    style_header(ws3, 2, len(bs_headers), "2E5FA3")

    for i, fac in enumerate(factors):
        rd = relation_boot[:, i]
        row_data = [fac,
                    round(rd.min(), 4), round(np.percentile(rd, 10), 4),
                    round(np.percentile(rd, 25), 4), round(np.median(rd), 4),
                    round(np.percentile(rd, 75), 4), round(np.percentile(rd, 90), 4),
                    round(rd.max(), 4), round(rd.mean(), 4), round(rd.std(), 4)]
        for ci, val in enumerate(row_data, 1):
            cell = ws3.cell(row=i+3, column=ci, value=val)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.font = Font(name="Arial", size=10)
            if ci > 1:
                if isinstance(val, float):
                    color = "DFF0D8" if val > 0 else "F2DEDE" if val < 0 else "FFFFFF"
                    cell.fill = PatternFill("solid", fgColor=color)

    add_thin_border(ws3, 2, 2+len(factors), 1, len(bs_headers))
    for ci, w in enumerate([14,8,8,8,10,8,8,8,10,8], 1):
        ws3.column_dimensions[get_column_letter(ci)].width = w

    wb.save(output_path)
    print(f"Excel results saved: {output_path}")

# Plot the Impact-Relation Map with bootstrap intervals.
def plot_irm(df, factors, prominence_boot, relation_boot, output_path, alpha):
    fig, axes = plt.subplots(1, 2, figsize=(18, 8))
    fig.patch.set_facecolor('#F8F9FA')

    # Chart 1: original IRM with bootstrap intervals.
    ax = axes[0]
    ax.set_facecolor('#FAFAFA')

    x = df["r+c (original)"].values
    y = df["r-c (original)"].values
    ci_col_lo = f"Bootstrap CI lower (r-c) {int((1-alpha)*100)}%"
    ci_col_hi = f"Bootstrap CI upper (r-c) {int((1-alpha)*100)}%"
    y_lo = df[ci_col_lo].values
    y_hi = df[ci_col_hi].values

    cause_color  = "#2196F3"
    effect_color = "#F44336"
    sig_marker   = "o"

    for i, fac in enumerate(df["Factor"].values):
        is_cause = df["Cause/Effect"].values[i] == "Cause"
        is_sig   = df["CI excludes zero?"].values[i] == "Stable"
        color = cause_color if is_cause else effect_color
        marker = sig_marker

        # Draw the interval bar.
        ax.plot([x[i], x[i]], [y_lo[i], y_hi[i]],
                color=color, alpha=0.3, linewidth=2, zorder=1)
        ax.plot([x[i]-0.005, x[i]+0.005], [y_lo[i], y_lo[i]],
                color=color, alpha=0.4, linewidth=1.5, zorder=1)
        ax.plot([x[i]-0.005, x[i]+0.005], [y_hi[i], y_hi[i]],
                color=color, alpha=0.4, linewidth=1.5, zorder=1)

        # Draw the estimate.
        ax.scatter(x[i], y[i], c=color, s=120 if is_sig else 70,
                   zorder=3, edgecolors='white', linewidth=1.5,
                   alpha=0.9, marker=marker if is_sig else "D")

        # Add the factor label.
        ax.annotate(fac, (x[i], y[i]),
                    textcoords="offset points", xytext=(6, 4),
                    fontsize=9, fontweight='bold', color=color,
                    bbox=dict(boxstyle="round,pad=0.2", fc="white", alpha=0.7, ec=color, lw=0.8))

    # Cause/effect dividing line.
    ax.axhline(0, color='#555555', linewidth=1.2, linestyle='--', alpha=0.7, zorder=2)

    # Mean prominence line.
    x_thresh = x.mean()
    ax.axvline(x_thresh, color='#9E9E9E', linewidth=1, linestyle=':', alpha=0.6, zorder=2)

    # Shade the four quadrants.
    y_lim = ax.get_ylim()
    x_lim = ax.get_xlim()
    ax.fill_between([x_lim[0], x_thresh], 0, y_lim[1], alpha=0.04, color='blue')
    ax.fill_between([x_thresh, x_lim[1]], 0, y_lim[1], alpha=0.06, color='blue')
    ax.fill_between([x_lim[0], x_thresh], y_lim[0], 0, alpha=0.04, color='red')
    ax.fill_between([x_thresh, x_lim[1]], y_lim[0], 0, alpha=0.06, color='red')

    ax.text(x_lim[1]*0.98, y_lim[1]*0.92, "CAUSE", ha='right', va='top',
            fontsize=11, color=cause_color, fontweight='bold', alpha=0.5)
    ax.text(x_lim[1]*0.98, y_lim[0]*0.92, "EFFECT", ha='right', va='bottom',
            fontsize=11, color=effect_color, fontweight='bold', alpha=0.5)

    ax.set_xlabel("r + c (Prominence)", fontsize=11, labelpad=8)
    ax.set_ylabel("r − c (Relation: Cause / Effect)", fontsize=11, labelpad=8)
    ax.set_title(f"Impact-Relation Map (IRM)\nBootstrap {int((1-alpha)*100)}% CI for r−c",
                 fontsize=12, fontweight='bold', pad=12)
    ax.grid(True, alpha=0.3, linestyle='--')

    legend_elements = [
        mpatches.Patch(color=cause_color, label='Cause'),
        mpatches.Patch(color=effect_color, label='Effect'),
        plt.scatter([], [], marker='o', c='gray', s=80, label='CI excludes zero'),
        plt.scatter([], [], marker='D', c='gray', s=60, label='CI includes zero'),
    ]
    ax.legend(handles=legend_elements, loc='upper left', fontsize=9,
              framealpha=0.8, edgecolor='#CCCCCC')

    # Chart 2: bootstrap violin plot for r-c.
    ax2 = axes[1]
    ax2.set_facecolor('#FAFAFA')

    fac_labels = df["Factor"].values
    n_fac = len(fac_labels)
    # Sort by descending original r-c.
    order = np.argsort(df["r-c (original)"].values)[::-1]
    boot_data = [relation_boot[:, df["Factor"].values.tolist().index(fac_labels[i])]
                 for i in order]
    labels_ord = [fac_labels[i] for i in order]
    orig_ord   = [df["r-c (original)"].values[i] for i in order]
    cause_ord  = [df["Cause/Effect"].values[i] for i in order]

    colors = [cause_color if c == "Cause" else effect_color for c in cause_ord]
    vp = ax2.violinplot(boot_data, positions=range(n_fac),
                        showmedians=True, showextrema=False)

    for i, (body, color) in enumerate(zip(vp['bodies'], colors)):
        body.set_facecolor(color)
        body.set_alpha(0.35)
        body.set_edgecolor(color)

    vp['cmedians'].set_color('#333333')
    vp['cmedians'].set_linewidth(2)

    # Draw the original estimates.
    ax2.scatter(range(n_fac), orig_ord, c=colors, s=60, zorder=5,
                edgecolors='white', linewidth=1.5)

    ax2.axhline(0, color='#555555', linewidth=1.2, linestyle='--', alpha=0.7)
    ax2.set_xticks(range(n_fac))
    ax2.set_xticklabels(labels_ord, rotation=45, ha='right', fontsize=9)
    ax2.set_ylabel("r − c (Bootstrap Distribution)", fontsize=11)
    ax2.set_title("Bootstrap distribution of r−c\n(Violin plot sorted by original r−c)",
                  fontsize=12, fontweight='bold', pad=12)
    ax2.grid(True, alpha=0.3, linestyle='--', axis='y')

    fig.suptitle("Z-Fuzzy DEMATEL Bootstrap Analysis", fontsize=14,
                 fontweight='bold', y=1.01, color='#1A3A6B')
    plt.tight_layout()
    plt.savefig(output_path, dpi=180, bbox_inches='tight',
                facecolor=fig.get_facecolor())
    plt.close()
    print(f"IRM chart saved: {output_path}")

# Command-line entry point.
def main():
    INPUT_FILE  = "INPUT.xlsx"
    OUTPUT_XLS  = "OUTPUT.xlsx"
    OUTPUT_IMG  = "OUTPUT.png"
    B     = 5000
    ALPHA = 0.05
    SEED  = 80

    print("=" * 60)
    print("  Bootstrap Z-Fuzzy DEMATEL Tool")
    print("=" * 60)

    # 1. Read input data.
    factors, experts = read_expert_sheets(INPUT_FILE)

    # 2. Bootstrap
    print(f"\nRunning bootstrap analysis (B={B:,})...")
    df, rc_plus, rc_minus, T_orig, p_boot, r_boot = run_bootstrap(
        factors, experts, B=B, seed=SEED, alpha=ALPHA)

    # 3. Export the Excel workbook.
    print("\nExporting Excel results...")
    export_excel(df, factors, T_orig, p_boot, r_boot, OUTPUT_XLS, B, ALPHA)

    # 4. Create the IRM chart.
    print("Creating the IRM chart...")
    plot_irm(df, factors, p_boot, r_boot, OUTPUT_IMG, ALPHA)

    # 5. Print the summary.
    print("\n" + "=" * 60)
    print("  ANALYSIS SUMMARY")
    print("=" * 60)
    display_cols = ["Factor", "r+c (original)", "r-c (original)", "Cause/Effect",
                    f"Bootstrap CI lower (r-c) {int((1-ALPHA)*100)}%",
                    f"Bootstrap CI upper (r-c) {int((1-ALPHA)*100)}%",
                    "Direction stability", "CI excludes zero?"]
    print(df[display_cols].to_string(index=False))
    print(f"\n  Experts: {len(experts)}")
    print(f"  Factors:       {len(factors)}")
    print(f"  Resamples:     {B:,}")
    print(f"  Confidence:    {int((1-ALPHA)*100)}%")
    print("=" * 60)

if __name__ == "__main__":
    main()

def run_pipeline(
    input_file,
    output_xls,
    output_img,
    B=5000,
    alpha=0.05,
    seed=80,
    start_row=2,
    start_col=2,
    n_rows=None,
    n_cols=None,
    header_row=1,
    fuzzy_scale=None,
    defuzz_func=None
):
    if fuzzy_scale is None:
        raise ValueError("Please import a fuzzy scale workbook.")

    if defuzz_func is None:
        defuzz_func = lambda l, m, u: (l + m + u)/3

    # Pass fuzzy settings to the workbook reader.
    factors, experts = read_expert_sheets(
        input_file,
        start_row=start_row,
        start_col=start_col,
        n_rows=n_rows,
        n_cols=n_cols,
        header_row=header_row,
        fuzzy_scale=fuzzy_scale,
        defuzz_func=defuzz_func
    )

    df, rc_plus, rc_minus, T_orig, p_boot, r_boot = run_bootstrap(
        factors, experts, B=B, seed=seed, alpha=alpha)

    export_excel(df, factors, T_orig, p_boot, r_boot, output_xls, B, alpha)
    plot_irm(df, factors, p_boot, r_boot, output_img, alpha)

    return df

if __name__ == "__main__":
    main()
