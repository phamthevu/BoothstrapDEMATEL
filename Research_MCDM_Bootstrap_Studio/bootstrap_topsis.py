"""Bootstrap TOPSIS for multiple expert rating sheets."""
from __future__ import annotations

import re
from pathlib import Path

import numpy as np
import pandas as pd
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt


def read_topsis_sheets(path, sheet_names=None, fuzzy_scale=None, defuzz_func=None):
    """Read expert sheets with alternatives in rows and criteria in columns.

    Layout: row 1 has criteria names from column B; column A has alternatives.
    Each expert sheet is named ``Ans N``. A single ``Data`` sheet is accepted
    for deterministic TOPSIS, but bootstrap requires at least two expert sheets.
    """
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = sheet_names or [s for s in wb.sheetnames if re.fullmatch(r"Ans\s*\d+", s, re.I)]
        names = sorted(names, key=lambda s: int(re.search(r"\d+", s).group())) if names and re.search(r"\d+", names[0]) else names
        if not names and "Data" in wb.sheetnames:
            names = ["Data"]
        if not names:
            raise ValueError("No Ans 1, Ans 2, ... expert worksheets or Data worksheet found")
        alternatives, criteria, matrices = None, None, []
        for sn in names:
            rows = list(wb[sn].iter_rows(values_only=True))
            if len(rows) < 3 or len(rows[0]) < 3:
                raise ValueError(f"Worksheet '{sn}' needs a header and at least two alternatives and two criteria")
            crit = [str(x).strip() for x in rows[0][1:] if x is not None and str(x).strip()]
            alts, matrix = [], []
            for row in rows[1:]:
                if not row or row[0] is None:
                    break
                alts.append(str(row[0]).strip())
                values = []
                for j, criterion in enumerate(crit, 1):
                    val = row[j] if j < len(row) else None
                    if val is None or str(val).strip() == "":
                        raise ValueError(f"Missing value in worksheet '{sn}', alternative '{alts[-1]}', criterion '{criterion}'")
                    try:
                        score = float(val)
                    except (TypeError, ValueError):
                        key = str(val).strip()
                        if fuzzy_scale is None or defuzz_func is None or key not in fuzzy_scale:
                            raise ValueError(f"Value '{val}' is not numeric and is not present in the fuzzy scale")
                        score = float(defuzz_func(*fuzzy_scale[key]))
                    if not np.isfinite(score):
                        raise ValueError(f"Non-finite value found in worksheet '{sn}'")
                    values.append(score)
                matrix.append(values)
            if alternatives is None:
                alternatives, criteria = alts, crit
            elif alts != alternatives or crit != criteria:
                raise ValueError(f"Alternative or criterion order in '{sn}' does not match the first worksheet")
            matrices.append(np.asarray(matrix, dtype=float))
        if len(criteria) < 2 or len(alternatives) < 2:
            raise ValueError("TOPSIS requires at least two alternatives and two criteria")
        return alternatives, criteria, matrices
    finally:
        wb.close()


def topsis_scores(decision_matrix, weights, benefit):
    x = np.asarray(decision_matrix, dtype=float)
    w = np.asarray(weights, dtype=float)
    benefit = np.asarray(benefit, dtype=bool)
    if x.ndim != 2 or x.shape[1] != len(w) or len(w) != len(benefit):
        raise ValueError("Matrix dimensions, weights, and criterion types do not match")
    if np.any(x < 0):
        raise ValueError("TOPSIS vector normalization currently requires non-negative values")
    denom = np.linalg.norm(x, axis=0)
    if np.any(denom == 0):
        raise ValueError("A criterion contains only zeros and cannot be normalized")
    v = (x / denom) * w
    ideal_best = np.where(benefit, v.max(axis=0), v.min(axis=0))
    ideal_worst = np.where(benefit, v.min(axis=0), v.max(axis=0))
    d_best = np.linalg.norm(v - ideal_best, axis=1)
    d_worst = np.linalg.norm(v - ideal_worst, axis=1)
    denom = d_best + d_worst
    return np.divide(d_worst, denom, out=np.full_like(d_worst, 0.5), where=denom != 0)


def run_bootstrap_topsis(alternatives, criteria, expert_matrices, B=2000,
                         alpha=.05, seed=42, weights=None, benefit=None):
    if len(expert_matrices) < 2:
        raise ValueError("Bootstrap TOPSIS requires at least two expert worksheets")
    if B < 100 or not 0 < alpha < 1:
        raise ValueError("Use at least 100 bootstrap samples and set alpha between 0 and 1")
    shape = (len(alternatives), len(criteria))
    if any(m.shape != shape for m in expert_matrices):
        raise ValueError("Expert matrix dimensions do not match the alternatives and criteria")
    w = np.ones(len(criteria), dtype=float) if weights is None else np.asarray(weights, dtype=float)
    if len(w) != len(criteria) or np.any(w < 0) or w.sum() <= 0:
        raise ValueError("Weights must be non-negative, one per criterion, and have a positive sum")
    w /= w.sum()
    benefit = np.ones(len(criteria), dtype=bool) if benefit is None else np.asarray(benefit, dtype=bool)
    if len(benefit) != len(criteria):
        raise ValueError("Choose Benefit or Cost for each criterion")
    point = topsis_scores(np.mean(expert_matrices, axis=0), w, benefit)
    rng = np.random.default_rng(seed)
    boot = np.empty((B, len(alternatives)), dtype=float)
    n_exp = len(expert_matrices)
    for b in range(B):
        ids = rng.integers(0, n_exp, n_exp)
        boot[b] = topsis_scores(np.mean([expert_matrices[i] for i in ids], axis=0), w, benefit)
    rank_boot = np.argsort(np.argsort(-boot, axis=1), axis=1) + 1
    lo, hi = alpha / 2, 1 - alpha / 2
    out = pd.DataFrame({
        "Alternative": alternatives,
        "TOPSIS score": point,
        "Score mean (BS)": boot.mean(axis=0),
        "Score SD (BS)": boot.std(axis=0, ddof=1),
        f"CI lower {int((1-alpha)*100)}%": np.quantile(boot, lo, axis=0),
        f"CI upper {int((1-alpha)*100)}%": np.quantile(boot, hi, axis=0),
        "Rank": pd.Series(point).rank(ascending=False, method="min").astype(int),
        "Mean rank (BS)": rank_boot.mean(axis=0),
        "Top rank frequency": (rank_boot == 1).mean(axis=0),
    }).sort_values("Rank").reset_index(drop=True)
    return out, boot


def export_topsis(df, boot, criteria, weights, alternatives, output_xlsx, output_png):
    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Bootstrap TOPSIS")
        pd.DataFrame({"Criterion": criteria, "Weight": weights}).to_excel(writer, index=False, sheet_name="Weights")
        pd.DataFrame(boot, columns=alternatives).to_excel(writer, index=False, sheet_name="Bootstrap scores")
        for ws in writer.book.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for col in ws.columns:
                width = min(34, max(12, max(len(str(c.value or "")) for c in col) + 2))
                ws.column_dimensions[col[0].column_letter].width = width
    fig, ax = plt.subplots(figsize=(max(8, len(df)*.8), 5.2))
    names = df["Alternative"].tolist()
    means = df["Score mean (BS)"].to_numpy()
    # CI column names are dynamic; locate them without assuming a confidence level.
    ci_lo = next(c for c in df.columns if c.startswith("CI lower"))
    ci_hi = next(c for c in df.columns if c.startswith("CI upper"))
    yerr = np.maximum(0, np.vstack([means - df[ci_lo].to_numpy(), df[ci_hi].to_numpy() - means]))
    ax.bar(names, means, color="#2E77D0", alpha=.82)
    ax.errorbar(names, means, yerr=yerr, fmt="none", ecolor="#26364A", capsize=4)
    ax.set_ylabel("Closeness coefficient")
    ax.set_title("Bootstrap TOPSIS: mean score and percentile CI")
    ax.tick_params(axis="x", rotation=30)
    ax.grid(axis="y", alpha=.25)
    fig.tight_layout()
    fig.savefig(output_png, dpi=180, bbox_inches="tight")
    plt.close(fig)
